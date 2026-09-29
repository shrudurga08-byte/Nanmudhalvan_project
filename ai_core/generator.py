"""Text cleaning and export formatters (.docx, .pdf, HTML preview)."""
import html
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fpdf import FPDF

from config import FOOTER_TEXT, LOGO_PATH

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2022": "-", "\u00a0": " ",
}


def sanitize_text(text: str) -> str:
    """Remove markdown leftovers and typographic characters."""
    text = text or ""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"^\s*#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"^\s*\*\s+", "- ", text, flags=re.M)
    text = text.replace("`", "")
    return text.strip()


def is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80:
        return False
    return bool(
        re.match(r"^\d+\.\s+[^.]{2,80}:?$", s)
        or s.endswith(":")
        or (s.isupper() and len(s) > 3)
    )


def is_bullet(line: str) -> bool:
    return line.strip().startswith(("- ", "* "))


def split_terms(terms: str | None) -> list[str]:
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


def _body_lines(text: str, doc_type: str) -> list[str]:
    lines = [l.rstrip() for l in sanitize_text(text).splitlines()]
    while lines and not lines[0].strip():
        lines.pop(0)
    if lines and lines[0].strip().lower() == (doc_type or "").strip().lower():
        lines.pop(0)  # title is added separately
    return lines


# ------------------------------------------------------------------ HTML
def format_html_preview(text: str) -> str:
    out = []
    for line in sanitize_text(text).splitlines():
        if not line.strip():
            continue
        safe = html.escape(line.strip())
        if is_bullet(line):
            out.append(f"<p style='margin:2px 0 2px 20px'>&bull; {html.escape(line.strip()[2:])}</p>")
        elif is_heading(line):
            out.append(f"<h4 style='margin:18px 0 4px 0;color:#e8e8e8'>{safe}</h4>")
        else:
            out.append(f"<p style='margin:6px 0;line-height:1.55'>{safe}</p>")
    return "".join(out)


# ------------------------------------------------------------------ DOCX
def format_docx(text: str, doc_type: str, terms: str | None = None) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    if LOGO_PATH.exists():
        doc.add_picture(str(LOGO_PATH), width=Inches(1.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type or "Legal Document")
    run.bold = True
    run.font.size = Pt(18)

    for line in _body_lines(text, doc_type):
        if not line.strip():
            continue
        if is_bullet(line):
            doc.add_paragraph(line.strip()[2:], style="List Bullet")
        elif is_heading(line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.add_run(line.strip()).bold = True
        else:
            p = doc.add_paragraph(line.strip())
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    items = split_terms(terms)
    if items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term / Condition"
        for c in table.rows[0].cells:
            for r in c.paragraphs[0].runs:
                r.bold = True
        for i, t in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(FOOTER_TEXT).italic = True

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
class _LegalPDF(FPDF):
    def header(self):
        if LOGO_PATH.exists():
            self.image(str(LOGO_PATH), x=(self.w - 40) / 2, y=8, w=40)
        self.set_y(30)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, f"{FOOTER_TEXT}  |  Page {self.page_no()}", align="C")


def _latin(s: str) -> str:
    return s.encode("latin-1", "replace").decode("latin-1")


def format_pdf(text: str, doc_type: str, terms: str | None = None) -> bytes:
    pdf = _LegalPDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 15)
    pdf.multi_cell(0, 9, _latin(doc_type or "Legal Document"), align="C",
                   new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    for line in _body_lines(text, doc_type):
        if not line.strip():
            continue
        s = _latin(line.strip())
        if is_bullet(line):
            pdf.set_font("Helvetica", "", 11)
            pdf.set_x(pdf.l_margin + 6)
            pdf.multi_cell(0, 6, "-  " + s[2:], new_x="LMARGIN", new_y="NEXT")
        elif is_heading(line):
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 7, s, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 6, s, align="J", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)

    items = split_terms(terms)
    if items:
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 7, "Summary of Key Terms", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for t in items:
            pdf.set_x(pdf.l_margin + 6)
            pdf.multi_cell(0, 6, "-  " + _latin(t), new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())
